"""Private, on-demand 300-question review workbook from the canonical V2 release.

Staff-only caller; do not persist this workbook in static assets or public Git.
It intentionally includes answer keys and reading rubrics for authorized review.
"""
from __future__ import annotations

from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


QUESTION_HEADERS = (
    "الترتيب", "القسم", "عنوان القسم", "معرف السؤال", "المرحلة",
    "التفاعل", "المهارة", "الدليل", "السؤال", "النص الظاهر للطالب",
    "الخيارات ومعرفاتها", "الإجابة الصحيحة", "الإجابات البديلة المعتمدة",
    "مرجع/معيار القراءة", "التلميح", "وسائط السؤال", "إصدار البنك",
)
SECTION_HEADERS = ("القسم", "النوع", "عنوان القسم", "الهدف", "التمهيد", "الأسئلة")


def _json(value: object) -> str:
    if value is None:
        return ""
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _safe_cell(value: object) -> object:
    """Escape spreadsheet formulas even in reviewer-visible authoring text."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if len(text) > 32767:
        raise ValueError("review bank cell exceeds XLSX text limit")
    if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def build_review_workbook(release: dict) -> bytes:
    """Faithfully export compiled bank data; never invent or normalize answers."""
    sections = [
        {"id": "pretest", "title": "الاختبار القبلي", "kind": "assessment"},
        *({"id": row["id"], "title": row["title"], "kind": "lesson",
           "objective": row.get("objective"), "intro": row.get("intro")} for row in release["lessons"]),
        *({"id": row["id"], "title": row["title"], "kind": "support",
           "objective": row.get("objective"), "intro": row.get("intro")} for row in release["supports"]),
        {"id": "posttest", "title": "الاختبار البعدي", "kind": "assessment"},
    ]
    by_id = {row["id"]: row for row in sections}
    order = {row["id"]: index for index, row in enumerate(sections)}
    ordered_questions = sorted(
        release["questions"],
        key=lambda question: (
            order[question.get("owner_id") or question["role"]],
            {"guided_practice": 0, "independent_check": 1, "support_check": 1}.get(question["role"], 0),
            question["id"],
        ),
    )
    workbook = Workbook()
    questions = workbook.active
    questions.title = "الأسئلة المعتمدة"
    questions.append(QUESTION_HEADERS)
    counts = {section_id: 0 for section_id in by_id}
    for index, question in enumerate(ordered_questions, 1):
        section_id = question.get("owner_id") or question["role"]
        counts[section_id] += 1
        answer = question["private_answer"]
        options = question.get("options", [])
        option_labels = {option["id"]: option.get("text") or _json(option.get("media", [])) for option in options}
        correct = [
            {"id": option_id, "text": option_labels.get(option_id, "")}
            for option_id in answer.get("correct_option_ids", [])
        ]
        stimulus = question.get("stimulus") or {}
        media_ids = [item["media_id"] for item in stimulus.get("media", [])]
        media_ids += [item["media_id"] for option in options for item in option.get("media", [])]
        row = [
            index, section_id, by_id[section_id]["title"], question["id"],
            question["role"], question["interaction"], question["skill_group"],
            question["evidence_type"], question["prompt"], stimulus.get("text"),
            _json(options), _json(correct), _json(answer.get("accepted_option_orders", [])),
            _json({"rubric": answer.get("rubric"), "reading_reference": question.get("reading_reference")}),
            question.get("hint"), _json(media_ids), release["version"],
        ]
        questions.append([_safe_cell(value) for value in row])

    sections_sheet = workbook.create_sheet("الأقسام والأهداف")
    sections_sheet.append(SECTION_HEADERS)
    for section in sections:
        intro = section.get("intro") or {}
        sections_sheet.append([
            _safe_cell(section["id"]), _safe_cell(section["kind"]),
            _safe_cell(section["title"]), _safe_cell(section.get("objective")),
            _safe_cell(intro.get("text")), counts[section["id"]],
        ])

    about = workbook.create_sheet("معلومات المراجعة")
    about.append(["المصدر", "بنك هِمّة V2 التنفيذي المترجم بواسطة compile_release"])
    about.append(["الإصدار", _safe_cell(release["version"])])
    about.append(["البصمة", _safe_cell(release["checksum"])])
    about.append(["عدد الأسئلة", len(ordered_questions)])
    about.append(["عدد الدروس", len(release["lessons"])])
    about.append(["عدد مسارات التقوية", len(release["supports"])])
    about.append(["خصوصية", "هذا الملف للمشرف والمراجع المخوّل فقط؛ يتضمن إجابات ومعايير خاصة. لا يُعرض للطفل ولا يُنشر كملف عام."])
    about.append(["الحالة", "مرجع المحتوى التنفيذي المعتمد في المستودع؛ لا يعني تصديره اعتماد النشر في بيئة الإنتاج."])

    for sheet in (questions, sections_sheet):
        sheet.sheet_view.rightToLeft = True
        sheet.freeze_panes = "D2" if sheet is questions else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="20364D")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = (
                55 if letter in ("I", "J", "K", "L", "M", "N") else 23
            )
            for cell in column:
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    about.sheet_view.rightToLeft = True
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 100

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
