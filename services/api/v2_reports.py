"""V2 persisted-evidence reports and private staff exports; never V1 measurements."""
from collections import defaultdict
from io import BytesIO
import json
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from db.learning_models import AdaptationDecision, AssessmentQuestionResponse, LessonSession, PlacementDecisionV2, SkillEvidence, StudentJourney, SupportAssignment, V2ReviewRetryTask
from db.models import AssessmentSession, AuditLog, ContentRelease, Student, User
from dependencies import get_current_user, get_db
from reports import _pdf_styles, _rtl
from v2_student_readmodel import _release_lessons

router = APIRouter(prefix='/v2/reports', tags=['v2-reports'])


def build_v2_reports(db: Session) -> dict:
    students = db.query(Student).order_by(Student.id).all()
    journeys = {row.student_id: row for row in db.query(StudentJourney).all()}
    releases = {row.id: row for row in db.query(ContentRelease).all()}
    selected = {}
    assessments = db.query(AssessmentSession).filter(AssessmentSession.v2_state.is_not(None)).order_by(AssessmentSession.id).all()
    official_ids = set()
    for row in assessments:
        if row.session_type in {'pretest', 'posttest'} and row.status == 'completed' and row.official_for_reporting and row.final_score is not None:
            selected[row.student_id, row.session_type] = row
    official_ids.update(row.id for row in selected.values())
    lessons = defaultdict(list)
    for row in db.query(LessonSession).all():
        lessons[row.student_id].append(row)
    supports = defaultdict(list)
    for row in db.query(SupportAssignment).all():
        supports[row.student_id].append(row)
    placements = {}
    for row in db.query(PlacementDecisionV2).order_by(PlacementDecisionV2.id).all():
        placements[row.student_id, row.content_release_id] = row.assigned_level
    # "resolved" only means that a support detour ended. Academic mastery
    # must be proven by independent checks and a persisted policy decision.
    mastered_support_ids = {
        assignment_id for assignment_id, in db.query(AdaptationDecision.support_assignment_id).filter(
            AdaptationDecision.reason_code.in_(
                ("support_check_resolved_gap", "alternate_check_resolved_gap"),
            ),
            AdaptationDecision.support_assignment_id.is_not(None),
        ).all()
    }
    unmastered_support_ids = {
        assignment_id for assignment_id, in db.query(AdaptationDecision.support_assignment_id).filter(
            AdaptationDecision.reason_code == "support_gap_unresolved_continue_without_mastery",
            AdaptationDecision.support_assignment_id.is_not(None),
        ).all()
    }
    mastered_support_ids.difference_update(unmastered_support_ids)
    pending = defaultdict(int)
    v2_session_owners = {row.id: row.student_id for row in assessments}
    requested_assessment_ids = db.query(V2ReviewRetryTask.assessment_response_id).filter(
        V2ReviewRetryTask.status == 'requested', V2ReviewRetryTask.assessment_response_id.is_not(None),
    )
    for session_id, in db.query(AssessmentQuestionResponse.assessment_session_id).filter(
        AssessmentQuestionResponse.status == 'pending_review',
        ~AssessmentQuestionResponse.id.in_(requested_assessment_ids),
    ).all():
        if session_id in v2_session_owners:
            pending[v2_session_owners[session_id]] += 1
    bank_cache = {}
    rows = []
    for student in students:
        journey = journeys.get(student.id)
        pre = selected.get((student.id, 'pretest'))
        post = selected.get((student.id, 'posttest'))
        # A retake or historical release is never compared across different banks.
        pair = bool(pre and post and pre.content_release_id == post.content_release_id)
        pre_score = float(pre.final_score) if pre else None
        post_score = float(post.final_score) if post else None
        improvement = round(post_score - pre_score, 2) if pair else None
        assigned = completed = skipped = percent = None
        support_rows = []
        if journey:
            if journey.content_release_id not in bank_cache:
                bank_cache[journey.content_release_id] = _release_lessons(db, journey.content_release_id)[0]
            bank = bank_cache[journey.content_release_id]
            student_lessons = [row for row in lessons[student.id] if row.content_release_id == journey.content_release_id]
            support_rows = [row for row in supports[student.id] if row.content_release_id == journey.content_release_id]
            if bank is not None and journey.placement_state in {'placed', 'posttest_ready', 'completed'}:
                done_ids = {row.lesson_id for row in student_lessons if row.status == 'completed'}
                starting_level = placements.get((student.id, journey.content_release_id))
                if starting_level is None:
                    starting_level = min([journey.current_level] + [int(row['level'][1]) for row in bank if row['id'] in done_ids])
                skipped_ids = {row['id'] for row in bank if int(row['level'][1]) < starting_level}
                assigned_ids = {row['id'] for row in bank} - skipped_ids
                assigned, skipped = len(assigned_ids), len(skipped_ids)
                completed = len(done_ids & assigned_ids)
                percent = round(100 * completed / assigned, 1) if assigned else None
        def score_record(session):
            if session is None:
                return None
            release = releases.get(session.content_release_id)
            return {'score': float(session.final_score), 'session_id': session.id, 'attempt_no': session.assessment_attempt_no,
                    'release_id': session.content_release_id, 'release_version': release.version if release else None}
        rows.append({'student_id': student.id, 'student_name': student.name,
            'account_status': 'active' if student.is_active else 'inactive',
            'current_level': journey.current_level if journey and journey.placement_state in {'placed','posttest_ready','completed'} else None,
            'journey_state': journey.placement_state if journey else 'not_started',
            'pretest': score_record(pre), 'posttest': score_record(post),
            'comparison_state': 'paired' if pair else 'release_mismatch' if pre and post else 'pending',
            'improvement_points': improvement,
            'relative_improvement': round(100 * improvement / pre_score, 2) if pair and pre_score else None,
            'progress': {'completed_lessons': completed, 'assigned_lessons': assigned, 'skipped_by_placement': skipped, 'percent': percent},
            'support': {'total': len(support_rows), 'resolved': sum(row.status == 'resolved' for row in support_rows),
                        'mastered': sum(row.id in mastered_support_ids for row in support_rows),
                        'needs_followup': sum(row.id in unmastered_support_ids for row in support_rows),
                        'waiting': sum(row.status in {'active','waiting_review'} for row in support_rows), 'escalated': sum(row.status == 'escalated' for row in support_rows)},
            'pending_assessment_reviews': pending[student.id]})
    skill_rows = defaultdict(lambda: {'graded': 0, 'correct': 0})
    # Only official pre/post responses and eligible graded V2 learning evidence.
    for row in db.query(AssessmentQuestionResponse).filter(AssessmentQuestionResponse.assessment_session_id.in_(official_ids),
            AssessmentQuestionResponse.status == 'graded', AssessmentQuestionResponse.is_correct.is_not(None)).all():
        key = (v2_session_owners[row.assessment_session_id], row.role, row.skill_group)
        skill_rows[key]['graded'] += 1
        skill_rows[key]['correct'] += int(row.is_correct is True)
    for row in db.query(SkillEvidence).filter(SkillEvidence.review_status == 'graded', SkillEvidence.counts_for_mastery.is_(True),
            SkillEvidence.is_correct.is_not(None)).all():
        journey = journeys.get(row.student_id)
        if journey is None or journey.content_release_id != row.content_release_id:
            continue
        key = (row.student_id, 'learning', row.skill_group)
        skill_rows[key]['graded'] += 1
        skill_rows[key]['correct'] += int(row.is_correct is True)
    skills = [{'student_id': student_id, 'phase': phase, 'skill_group': skill, **values,
               'observed_accuracy_percent': round(100 * values['correct'] / values['graded'], 2)}
              for (student_id, phase, skill), values in sorted(skill_rows.items())]
    def mean(values):
        return round(sum(values) / len(values), 2) if values else None
    return {'source': 'himma_v2_persisted_evidence', 'students': rows, 'skills': skills,
        'summary': {'students': len(rows), 'official_pretests': sum(row['pretest'] is not None for row in rows),
                    'official_posttests': sum(row['posttest'] is not None for row in rows),
                    'paired': sum(row['comparison_state'] == 'paired' for row in rows),
                    'average_pretest': mean([row['pretest']['score'] for row in rows if row['pretest']]),
                    'average_posttest': mean([row['posttest']['score'] for row in rows if row['posttest']]),
                    'average_improvement_points': mean([row['improvement_points'] for row in rows if row['improvement_points'] is not None])},
        'methodology': 'V2 official completed pre/post only; same-release pairs; graded eligible lesson evidence; support mastered only on successful independent-check decisions (not activity closure); unresolved gaps need follow-up; null is not zero.'}


@router.get('/summary')
def summary(_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return build_v2_reports(db)


def _safe_cell(value):
    if isinstance(value, str) and value.lstrip(' \t\r\n').startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def xlsx_bytes(report: dict, *, skills: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'مهارات V2' if skills else 'نتائج V2'
    sheet.sheet_view.rightToLeft = True
    if skills:
        names = {row['student_id']: row['student_name'] for row in report['students']}
        sheet.append(['المعرف', 'الطالب', 'مرحلة الدليل', 'المهارة', 'استجابات مقيمة', 'صحيح', 'دقة مرصودة %'])
        for row in report['skills']:
            sheet.append([row['student_id'], _safe_cell(names[row['student_id']]), row['phase'], row['skill_group'], row['graded'], row['correct'], row['observed_accuracy_percent']])
    else:
        sheet.append(['المعرف', 'الطالب', 'المستوى V2', 'القبلي الرسمي', 'البعدي الرسمي', 'التحسن بالنقاط', 'التحسن النسبي %', 'حالة المقارنة', 'دروس مكتملة', 'دروس معينة', 'متجاوزة بالتسكين', 'دعم مثبت الإتقان', 'دعم كلي', 'قراءات معلقة', 'إصدار القبلي', 'إصدار البعدي'])
        for row in report['students']:
            sheet.append([row['student_id'], _safe_cell(row['student_name']), row['current_level'],
                row['pretest']['score'] if row['pretest'] else None, row['posttest']['score'] if row['posttest'] else None,
                row['improvement_points'], row['relative_improvement'], row['comparison_state'],
                row['progress']['completed_lessons'], row['progress']['assigned_lessons'], row['progress']['skipped_by_placement'],
                row['support']['mastered'], row['support']['total'], row['pending_assessment_reviews'],
                row['pretest']['release_version'] if row['pretest'] else None, row['posttest']['release_version'] if row['posttest'] else None])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='20364D')
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
        for cell in column:
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    notes = workbook.create_sheet('مصدر القياس')
    notes.sheet_view.rightToLeft = True
    notes.append(['المصدر', report['source']]); notes.append(['السياسة', report['methodology']])
    buffer = BytesIO(); workbook.save(buffer)
    return buffer.getvalue()


def pdf_bytes(report: dict) -> bytes:
    font, title, body, small = _pdf_styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm, bottomMargin=12*mm)
    def paragraph(value, style=small, width=None):
        text = str(value if value is not None else '—')
        lines, current = [], ''
        for word in text.split():
            candidate = f'{current} {word}'.strip()
            if width and current and stringWidth(_rtl(candidate), font, style.fontSize) > width - 12:
                lines.append(current); current = word
            else:
                current = candidate
        lines.append(current)
        return Paragraph('<br/>'.join(escape(_rtl(line)) for line in lines), style)
    story = [paragraph('تقارير هِمّة V2 — النتائج المعتمدة', title), Spacer(1, 5*mm),
             paragraph('النتائج من الاختبارات المكتملة المعتمدة فقط. العلامة — تعني أن النتيجة غير متاحة، ولا تعني صفرًا.', body), Spacer(1, 4*mm)]
    headers = ['الطالب', 'المستوى', 'قبلي', 'بعدي', 'التحسن بالنقاط', 'دروس مكتملة/معينة', 'دعم متحقق/كلي']
    widths = [58*mm, 22*mm, 24*mm, 24*mm, 32*mm, 48*mm, 30*mm]
    table = [[paragraph(value, width=width) for value, width in zip(headers, widths)][::-1]]
    for row in report['students']:
        done, assigned = row['progress']['completed_lessons'], row['progress']['assigned_lessons']
        values = [row['student_name'], row['current_level'], row['pretest']['score'] if row['pretest'] else None,
                  row['posttest']['score'] if row['posttest'] else None, row['improvement_points'],
                  f'{done}/{assigned}' if assigned is not None else None, f"{row['support']['mastered']}/{row['support']['total']}"]
        table.append([paragraph(value, width=width) for value, width in zip(values, widths)][::-1])
    if len(table) == 1:
        story.append(paragraph('لا توجد حسابات في العينة.', body))
    else:
        grid = Table(table, colWidths=widths[::-1], repeatRows=1, hAlign='RIGHT')
        grid.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),font), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE8F2')),
            ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#DCE8F2')), ('VALIGN',(0,0),(-1,-1),'TOP'), ('TOPPADDING',(0,0),(-1,-1),7), ('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.append(grid)
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font, 8)
        canvas.drawCentredString(landscape(A4)[0] / 2, 7*mm, _rtl(f'هِمّة V2 - صفحة {document.page}'))
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


@router.get('/exports/{filename}')
def export_report(filename: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if filename not in {'cohort.xlsx', 'cohort.pdf', 'skills.xlsx'}:
        raise HTTPException(status_code=404, detail='ملف التقرير غير موجود')
    report = build_v2_reports(db)
    payload = pdf_bytes(report) if filename.endswith('.pdf') else xlsx_bytes(report, skills=filename == 'skills.xlsx')
    db.add(AuditLog(actor_role='researcher', actor_id=user.id, action='v2.report.exported', entity_type='v2_report', entity_id=filename,
                    details=json.dumps({'source': report['source'], 'format': filename.split('.')[-1]})))
    db.commit()
    return StreamingResponse(BytesIO(payload), media_type='application/pdf' if filename.endswith('.pdf') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition': f'attachment; filename="himma-v2-{filename}"', 'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff'})


@router.get('/students/{student_id}/export.pdf')
def export_student_report(student_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    report = build_v2_reports(db)
    report['students'] = [row for row in report['students'] if row['student_id'] == student_id]
    if not report['students']:
        raise HTTPException(status_code=404, detail='الطالب غير موجود')
    payload = pdf_bytes(report)
    db.add(AuditLog(actor_role='researcher', actor_id=user.id, action='v2.report.exported', entity_type='student',
                    entity_id=str(student_id), details=json.dumps({'source': report['source'], 'format': 'pdf'})))
    db.commit()
    return StreamingResponse(BytesIO(payload), media_type='application/pdf',
        headers={'Content-Disposition': f'attachment; filename="himma-v2-student-{student_id}.pdf"',
                 'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff'})
