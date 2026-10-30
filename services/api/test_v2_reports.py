from io import BytesIO

from openpyxl import load_workbook

from conftest import TestingSessionLocal
from db.models import AssessmentSession, AuditLog, ContentRelease, Student
from v2_content.importer import register_candidate_release


def _official_pair(pre=0, post=50, *, same_release=True, name='اصطناعي تقرير'):
    with TestingSessionLocal() as db:
        release = register_candidate_release(db)
        student = db.query(Student).one()
        student.name = name
        other = release
        if not same_release:
            other = ContentRelease(version='other-report-bank', checksum='d'*64)
            db.add(other); db.flush()
        db.add_all([
            AssessmentSession(student_id=student.id, session_type='pretest', status='completed', v2_state='completed',
                content_release_id=release.id, official_for_reporting=True, final_score=pre),
            AssessmentSession(student_id=student.id, session_type='posttest', status='completed', v2_state='completed',
                content_release_id=other.id, official_for_reporting=True, final_score=post),
            AssessmentSession(student_id=student.id, session_type='pretest', status='completed', final_score=99,
                assessment_attempt_no=2, official_for_reporting=False),
            AssessmentSession(student_id=student.id, session_type='posttest', status='completed', v2_state='completed',
                content_release_id=release.id, assessment_attempt_no=2, official_for_reporting=False, final_score=100),
        ])
        db.commit()
        return student.id


def test_v2_reports_preserve_zero_and_only_compare_official_same_release_evidence(researcher_client):
    _official_pair()
    result = researcher_client.get('/v2/reports/summary')
    assert result.status_code == 200, result.text
    data = result.json(); row = data['students'][0]
    assert data['source'] == 'himma_v2_persisted_evidence'
    assert row['pretest']['score'] == 0 and row['posttest']['score'] == 50
    assert row['improvement_points'] == 50 and row['relative_improvement'] is None
    assert data['summary']['paired'] == 1
    assert data['summary']['average_pretest'] == 0
    assert data['summary']['average_improvement_points'] == 50
    assert row['progress']['assigned_lessons'] is None
    assert 'access_code' not in result.text and 'private_answer' not in result.text


def test_different_releases_are_not_reported_as_improvement(researcher_client):
    _official_pair(40, 80, same_release=False)
    data = researcher_client.get('/v2/reports/summary').json()
    assert data['students'][0]['comparison_state'] == 'release_mismatch'
    assert data['students'][0]['improvement_points'] is None
    assert data['summary']['paired'] == 0 and data['summary']['average_improvement_points'] is None


def test_student_cannot_read_or_export_v2_cohort(student_client):
    for path in ['/v2/reports/summary', '/v2/reports/exports/cohort.xlsx', '/v2/reports/exports/cohort.pdf', '/v2/reports/students/1/export.pdf']:
        assert student_client.get(path).status_code == 403


def test_v2_xlsx_contains_official_numbers_and_neutralizes_formula_names(researcher_client):
    _official_pair(name='=HYPERLINK("https://example.invalid")')
    response = researcher_client.get('/v2/reports/exports/cohort.xlsx')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.worksheets[0]
    assert sheet['B2'].value.startswith("'=HYPERLINK") and sheet['B2'].data_type == 's'
    assert sheet['D2'].value == 0 and sheet['E2'].value == 50 and sheet['F2'].value == 50
    assert sheet['G2'].value is None
    assert sheet.sheet_view.rightToLeft is True
    assert response.headers['cache-control'] == 'private, no-store'
    with TestingSessionLocal() as db:
        assert db.query(AuditLog).filter(AuditLog.action == 'v2.report.exported').count() == 1
        assert db.query(Student).one().name.startswith('=HYPERLINK')


def test_v2_pdf_and_individual_export_are_real_private_documents(researcher_client):
    student_id = _official_pair(name='اصطناعي <اختبار> & تقرير')
    for path in ['/v2/reports/exports/cohort.pdf', f'/v2/reports/students/{student_id}/export.pdf']:
        response = researcher_client.get(path)
        assert response.status_code == 200, response.text[:100]
        assert response.content.startswith(b'%PDF-') and len(response.content) > 2000
        assert response.headers['cache-control'] == 'private, no-store'
    assert researcher_client.get('/v2/reports/students/999999/export.pdf').status_code == 404


def test_unapproved_or_pending_evidence_does_not_create_a_score(researcher_client):
    with TestingSessionLocal() as db:
        student = db.query(Student).one()
        db.add(AssessmentSession(student_id=student.id, session_type='pretest', status='in_progress',
                                 v2_state='waiting_review', final_score=100, official_for_reporting=False))
        db.commit()
    data = researcher_client.get('/v2/reports/summary').json()
    assert data['students'][0]['pretest'] is None
    assert data['summary']['average_pretest'] is None and data['skills'] == []


def test_skill_summary_and_progress_keep_only_eligible_v2_evidence(researcher_client):
    from db.learning_models import LessonQuestionResponse, LessonSession, SkillEvidence, StudentJourney
    from v2_content.repository import get_release
    student_id = _official_pair()
    with TestingSessionLocal() as db:
        release = register_candidate_release(db)
        db.add(StudentJourney(student_id=student_id, content_release_id=release.id,
                              current_level=2, current_lesson_id='L2-01', placement_state='placed'))
        session = LessonSession(student_id=student_id, content_release_id=release.id,
                                lesson_id='L1-01', stage='complete', status='completed')
        db.add(session); db.flush()
        questions = [q for q in get_release()['questions'] if q['owner_id'] == 'L1-01' and q['role'] == 'independent_check'][:2]
        for index, question in enumerate(questions):
            row = LessonQuestionResponse(lesson_session_id=session.id, question_id=question['id'],
                role='independent_check', submit_key=f'evidence-{index}', answer_json={},
                status='graded' if index == 0 else 'pending_review', is_correct=True if index == 0 else None)
            db.add(row); db.flush()
            db.add(SkillEvidence(response_id=row.id, lesson_session_id=session.id, student_id=student_id,
                content_release_id=release.id, question_id=question['id'], source_kind='lesson', role='independent_check',
                skill_group=question['skill_group'], evidence_type=question['evidence_type'],
                review_status='graded' if index == 0 else 'pending_review', is_correct=True if index == 0 else None,
                counts_for_mastery=True, hint_used=False, policy_version='synthetic-test'))
        db.commit()
    report = researcher_client.get('/v2/reports/summary').json()
    home = researcher_client.get(f'/v2/student/admin/students/{student_id}/home').json()
    assert report['students'][0]['progress']['assigned_lessons'] == home['progress']['assigned_lessons'] == 19
    assert report['students'][0]['progress']['completed_lessons'] == home['progress']['completed_lessons'] == 1
    assert len(report['skills']) == 1
    assert report['skills'][0]['graded'] == report['skills'][0]['correct'] == 1
    assert report['skills'][0]['phase'] == 'learning'
    assert report['skills'][0]['observed_accuracy_percent'] == 100


def test_bundled_arabic_font_works_without_host_font_packages(monkeypatch):
    from pathlib import Path
    from reports import _find_report_font
    bundled = Path(__file__).resolve().parents[2] / 'assets/fonts/DejaVuSans.ttf'
    monkeypatch.delenv('HIMMA_REPORT_FONT_PATH', raising=False)
    monkeypatch.setattr(Path, 'is_file', lambda path: str(path) == str(bundled))
    assert _find_report_font() == bundled
    assert bundled.stat().st_size > 100000


def test_support_report_requires_independent_mastery_not_merely_resolved_status(researcher_client):
    from db.learning_models import AdaptationDecision, LessonSession, SupportAssignment
    student_id = _official_pair()
    with TestingSessionLocal() as db:
        release = db.query(ContentRelease).filter(ContentRelease.version.is_not(None)).order_by(ContentRelease.id).first()
        assert release is not None
        lesson = LessonSession(student_id=student_id, content_release_id=release.id,
                               lesson_id="L1-01", stage="outcome", status="completed")
        db.add(lesson)
        db.flush()
        for index, reason in enumerate(("support_check_resolved_gap",
                                        "support_gap_unresolved_continue_without_mastery"), 1):
            assignment = SupportAssignment(
                student_id=student_id, lesson_session_id=lesson.id,
                content_release_id=release.id, support_id="SUP-L1-01",
                focus=f"focus-{index}", target_skill="S1",
                source_lesson_id="L1-01", source_level=1, support_level=1,
                mode="targeted_support", status="resolved", verification_round=0,
                practice_question_ids=[], check_question_ids=[],
                return_to={}, policy_version="synthetic-v2",
            )
            db.add(assignment)
            db.flush()
            db.add(AdaptationDecision(
                student_id=student_id, lesson_session_id=lesson.id,
                content_release_id=release.id, support_assignment_id=assignment.id,
                decision_key=f"support-test-{index}", decision_kind="next_lesson",
                reason_code=reason, evidence_ids=[], explanation_json={},
                policy_version="synthetic-v2",
            ))
        db.commit()
    response = researcher_client.get("/v2/reports/summary")
    assert response.status_code == 200, response.text
    support = response.json()["students"][0]["support"]
    assert support["total"] == 2 and support["resolved"] == 2
    assert support["mastered"] == 1 and support["needs_followup"] == 1
    payload = researcher_client.get("/v2/reports/exports/cohort.xlsx")
    assert payload.status_code == 200, payload.text[:100]
    sheet = load_workbook(BytesIO(payload.content)).worksheets[0]
    assert "دعم مثبت الإتقان" in [cell.value for cell in sheet[1]]
    assert sheet["L2"].value == 1 and sheet["M2"].value == 2


def test_requested_rerecord_is_not_reported_as_staff_pending_review(researcher_client):
    from db.learning_models import AssessmentQuestionResponse, V2ReviewRetryTask
    from db.models import User
    with TestingSessionLocal() as db:
        student = db.query(Student).one()
        reviewer = db.query(User).first()
        assert reviewer is not None
        session = AssessmentSession(
            student_id=student.id, session_type="pretest", status="in_progress",
            v2_state="waiting_review", official_for_reporting=False,
        )
        db.add(session); db.flush()
        response = AssessmentQuestionResponse(
            assessment_session_id=session.id, question_id="synthetic-report-read",
            role="pretest", status="pending_review", is_correct=None,
            answer_json={}, submit_key="report-retry-only",
            skill_group="S1", evidence_type="read_aloud",
        )
        db.add(response); db.flush()
        retry = V2ReviewRetryTask(
            student_id=student.id, requested_by=reviewer.id,
            assessment_response_id=response.id,
            status="requested", request_kind="rerecord", reason="قراءة تحتاج إعادة",
        )
        db.add(retry); db.commit()
    report = researcher_client.get("/v2/reports/summary")
    assert report.status_code == 200, report.text
    assert report.json()["students"][0]["pending_assessment_reviews"] == 0
